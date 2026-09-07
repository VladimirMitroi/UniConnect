import pika
import threading
import fitz  # PyMuPDF
from docx import Document
from pptx import Presentation
import io
import fitz  # PyMuPDF
from fastapi import FastAPI, HTTPException
from pydantic import BaseModel
from minio import Minio
import google.generativeai as genai
import json
import re
import uvicorn
from gtts import gTTS
import uuid

genai.configure(api_key="AIzaSyAp7CspUhuK53ZkBbIWB4QNULCi_DFo1fQ") 
model = genai.GenerativeModel('gemini-2.5-flash')

minio_client = Minio(
    "localhost:9000",
    access_key="minioadmin",
    secret_key="minioadmin123",
    secure=False
)

app = FastAPI(title="UniConnect AI Service")

def rabbitmq_consumer():
    connection = pika.BlockingConnection(pika.ConnectionParameters(host='localhost'))
    channel = connection.channel()
    channel.queue_declare(queue='course_queue', durable=True)

    def callback(ch, method, properties, body):
        body_str = body.decode()
        
        nume_fisier = body_str
        numar_intrebari = 3
        tip_cerut = "single"
        
        try:
            mesaj_primit = json.loads(body_str)
            nume_fisier = mesaj_primit.get("fileName", body_str)
            numar_intrebari = mesaj_primit.get("numQuestions", 3)
            tip_cerut = mesaj_primit.get("questionType", "single")
            test_id = mesaj_primit.get("testId")
        except json.JSONDecodeError:
            pass 
            
        print(f"\n[📥 NOTIFICARE] Descarc: {nume_fisier} | Generez: {numar_intrebari} întrebări | Tip: {tip_cerut}")
        
        try:
            response = minio_client.get_object("uniconnect-documents", nume_fisier)
            pdf_content = response.read()
            response.close()
            response.release_conn()
            
            pdf_document = fitz.open(stream=pdf_content, filetype="pdf")
            text_extras = ""
            for page_num in range(len(pdf_document)):
                text_extras += pdf_document.load_page(page_num).get_text()
            pdf_document.close()
            
            print(f"[✅ TEXT EXTRAS] {len(text_extras)} caractere. Construiesc promptul...")

            if tip_cerut == "single":
                instructiune_tip = "Toate întrebările trebuie să aibă EXACT UN SINGUR răspuns corect și type: 'single'."
            elif tip_cerut == "multiple":
                instructiune_tip = "Toate întrebările trebuie să aibă DOUĂ SAU MAI MULTE răspunsuri corecte și type: 'multiple'."
            else:
                instructiune_tip = "Generează un amestec. Pentru cele cu un singur răspuns pune type: 'single', iar pentru cele cu mai multe răspunsuri pune type: 'multiple'."

            prompt = f"""
            Ești un profesor universitar extrem de riguros. Analizează textul cursului de mai jos.
            
            SARCINĂ OBLIGATORIE:
            Generează EXACT {numar_intrebari} întrebări grilă. Niciuna în plus, niciuna în minus!
            
            REGULI STRICTE:
            1. LIMBA: Absolut toate întrebările, variantele de răspuns și explicațiile TREBUIE să fie scrise în limba ROMÂNĂ.
            2. FORMAT: Returnează STRICT un array JSON valid, conform structurii de mai jos.
            3. REGULĂ DE AUR: {instructiune_tip}
            
            IMPORTANT: 
            - Dacă type este 'single', lista 'correctAnswers' trebuie să aibă EXACT UN element.
            - Dacă type este 'multiple', lista 'correctAnswers' trebuie să aibă MINIM DOUĂ elemente.
            - NU folosi niciodată valoarea 'mix' în câmpul 'type' din JSON. Folosește doar 'single' sau 'multiple'.
            
            Structură JSON necesară:
            [
              {{
                "question": "Textul întrebării în română?",
                "options": ["Varianta 1", "Varianta 2", "Varianta 3", "Varianta 4"],
                "correctAnswers": ["Varianta 1", "Varianta 3"], 
                "type": "multiple" 
              }}
            ]

            Textul cursului:
            {text_extras}
            """

            raspuns_ai = model.generate_content(prompt)
            
            text_gemini = raspuns_ai.text
            
            print("\n[🔍 Răspuns brut de la Gemini (primele 200 caractere)]:")
            print(text_gemini[:200]) 
            
            match = re.search(r'\[.*\]', text_gemini, re.DOTALL)
            
            if match:
                raw_text = match.group(0)
            else:
                raw_text = text_gemini.strip() 
                
            try:
                date_json = json.loads(raw_text)

                limita = int(numar_intrebari) 
            
                date_json = date_json[:limita] 
            
                for intrebare in date_json:
                    intrebare["testId"] = test_id 
                
                mesaj_final_pentru_java = json.dumps(date_json)
            except json.JSONDecodeError as e:
                print(f"[❌ EROARE CRITICĂ] AI-ul nu a generat JSON valid. Text extras: {raw_text}")
                raise e 
            
            print("\n[🧠 REZULTAT AI CURĂȚAT ȘI MODIFICAT]")
            print(mesaj_final_pentru_java[:300] + "... [TRUNCHIAT]") 
            print("-" * 50)
            
            channel.queue_declare(queue='results_queue', durable=True)
            channel.basic_publish(
                exchange='',
                routing_key='results_queue',
                body=mesaj_final_pentru_java
            )
            print("[📤 SUCCES] Am trimis JSON-ul către Java!\n")

        except Exception as e:
            print(f"[❌ EROARE] Ceva a mers prost la procesarea {nume_fisier}: {e}")
        
    channel.basic_consume(queue='course_queue', on_message_callback=callback, auto_ack=True)
    print(" [*] RabbitMQ Consumer rulează silențios în fundal...")
    channel.start_consuming()

@app.on_event("startup")
def startup_event():
    thread = threading.Thread(target=rabbitmq_consumer, daemon=True)
    thread.start()

@app.get("/")
def health_check():
    return {"status": "online", "service": "UniConnect AI Worker"}

class ChatRequest(BaseModel):
    message: str
    fileNames: list[str]

class AutoGradeRequest(BaseModel):
    assignment_description: str
    file_name: str

class PodcastRequest(BaseModel):
    file_name: str

@app.post("/chat")
def chat_with_course(req: ChatRequest):
    print(f"\n[💬 CHATBOT] Mesaj primit: {req.message}")
    print(f"[📚 CONTEXT] Caut în fișierele: {req.fileNames}")
    
    context_text = ""
    
    for file_name in req.fileNames:
        ext = file_name.split('.')[-1].lower() if '.' in file_name else ''
        if ext not in ['pdf', 'docx', 'txt', 'pptx']:
            continue 
            
        try:
            response = minio_client.get_object("uniconnect-documents", file_name)
            file_content = response.read()
            response.close()
            response.release_conn()
            
            if ext == 'pdf':
                pdf_document = fitz.open(stream=file_content, filetype="pdf")
                for page_num in range(len(pdf_document)):
                    context_text += pdf_document.load_page(page_num).get_text() + "\n"
                pdf_document.close()
            elif ext == 'txt':
                context_text += file_content.decode('utf-8') + "\n"
            elif ext == 'docx':
                doc = Document(io.BytesIO(file_content))
                for para in doc.paragraphs:
                    context_text += para.text + "\n"
            elif ext == 'pptx':
                prs = Presentation(io.BytesIO(file_content))
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            context_text += shape.text + "\n"

            print(f"[✅ FIȘIER CITIT] {file_name}")
        except Exception as e:
            print(f"[⚠️ AVERTISMENT] Nu am putut citi {file_name}: {e}")
            
    if not context_text.strip():
         return {"response": "Nu am găsit materiale compatibile (PDF, DOCX, TXT, PPTX) în acest curs din care să pot învăța."}
         
    context_text = context_text[:200000]

    prompt = f"""
    Ești asistentul universitar virtual (AI) pentru acest curs pe platforma UniConnect.
    Răspunde la întrebarea studentului folosind STRICT informațiile din documentele cursului de mai jos.
    Dacă întrebarea nu are nicio legătură cu conținutul cursului sau informația nu se regăsește în text, 
    spune politicos că nu deții aceste informații în materialele primite, fără a inventa.
    Răspunde clar, concis, cu bullet points dacă este cazul, și mereu în limba română.

    MATERIALE CURS:
    ---
    {context_text}
    ---
    
    ÎNTREBAREA STUDENTULUI:
    {req.message}
    """
    
    try:
        raspuns_ai = model.generate_content(prompt)
        return {"response": raspuns_ai.text}
    except Exception as e:
        print(f"[❌ EROARE GEMINI] {e}")
        raise HTTPException(status_code=500, detail="Eroare la generarea răspunsului AI.")

@app.post("/flashcards")
def generate_flashcards(req: ChatRequest):
    context_text = ""
    for file_name in req.fileNames:
        ext = file_name.split('.')[-1].lower() if '.' in file_name else ''
        if ext not in ['pdf', 'docx', 'txt', 'pptx']:
            continue
            
        try:
            response = minio_client.get_object("uniconnect-documents", file_name)
            file_content = response.read()
            response.close()
            response.release_conn()
            
            if ext == 'pdf':
                pdf_document = fitz.open(stream=file_content, filetype="pdf")
                for page_num in range(len(pdf_document)):
                    context_text += pdf_document.load_page(page_num).get_text() + "\n"
                pdf_document.close()
            elif ext == 'txt':
                context_text += file_content.decode('utf-8') + "\n"
            elif ext == 'docx':
                doc = Document(io.BytesIO(file_content))
                for para in doc.paragraphs:
                    context_text += para.text + "\n"
            elif ext == 'pptx':
                prs = Presentation(io.BytesIO(file_content))
                for slide in prs.slides:
                    for shape in slide.shapes:
                        if hasattr(shape, "text"):
                            context_text += shape.text + "\n"
        except Exception as e:
            print(f"[⚠️ AVERTISMENT] Nu am putut citi {file_name}: {e}")
            
    if not context_text.strip():
         return {"flashcards": []}
         
    context_text = context_text[:200000]
    
    prompt = f"""Ești un expert în educație. Extrage cele mai importante concepte din următorul text de curs și creează o listă de maxim 10 Flashcards (cartonașe de memorare).
    Textul cursului:
    {context_text}
    
    Returnează STRICT un JSON valid cu următoarea structură, fără absolut nimic altceva:
    {{
       "flashcards": [
          {{
             "concept": "Numele conceptului / Întrebarea",
             "definition": "Definiția sau răspunsul clar și concis (max 2 propoziții)"
          }}
       ]
    }}
    """
    
    try:
        response = model.generate_content(prompt)
        raw_text = response.text.strip()
        
        if raw_text.startswith("```json"):
            raw_text = raw_text.split("```json")[1]
        if raw_text.endswith("```"):
            raw_text = raw_text.rsplit("```", 1)[0]
            
        raw_text = raw_text.strip()
        import json
        result = json.loads(raw_text)
        return result
    except Exception as e:
        print(f"[❌ EROARE FLASHCARDS] {e}")
        return {"flashcards": []}

@app.post("/grade-submission")
def grade_submission(req: AutoGradeRequest):
    context_text = ""
    file_name = req.file_name
    ext = file_name.split('.')[-1].lower() if '.' in file_name else ''
    
    if ext not in ['pdf', 'docx', 'txt', 'pptx']:
         return {"suggestedGrade": 0, "feedback": "Fișierul nu poate fi citit de AI. Corectează manual."}
         
    try:
        response = minio_client.get_object("uniconnect-documents", file_name)
        file_content = response.read()
        response.close()
        response.release_conn()
        
        if ext == 'pdf':
            pdf_document = fitz.open(stream=file_content, filetype="pdf")
            for page_num in range(len(pdf_document)):
                context_text += pdf_document.load_page(page_num).get_text() + "\n"
            pdf_document.close()
        elif ext == 'txt':
            context_text += file_content.decode('utf-8') + "\n"
        elif ext == 'docx':
            doc = Document(io.BytesIO(file_content))
            for para in doc.paragraphs:
                context_text += para.text + "\n"
        elif ext == 'pptx':
            prs = Presentation(io.BytesIO(file_content))
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        context_text += shape.text + "\n"
    except Exception as e:
        print(f"[⚠️ AVERTISMENT] Nu am putut citi {file_name}: {e}")
        return {"suggestedGrade": 0, "feedback": f"Eroare la citirea fișierului: {str(e)}"}
        
    if not context_text.strip():
         return {"suggestedGrade": 0, "feedback": "Fișierul pare gol sau textul nu poate fi extras."}
         
    context_text = context_text[:150000] 
    
    prompt = f"""Ești un asistent universitar. Analizează lucrarea trimisă de student și acordă o notă de la 1 la 10 pe baza cerinței profesorului. Fii obiectiv și critic, dar constructiv. Dacă cerința nu are criterii clare, judecă după calitatea, coerența și corectitudinea materialului. Dacă lucrarea e pe lângă subiect, pune notă mică.
    
    CERINȚA TEMEI:
    {req.assignment_description}
    
    LUCRAREA STUDENTULUI:
    {context_text}
    
    Returnează STRICT un JSON valid cu următoarea structură:
    {{
       "suggestedGrade": nota (număr de la 1 la 10, poate fi și cu zecimale gen 8.5),
       "feedback": "Scurt feedback către student care să explice nota. Maxim 3 fraze."
    }}
    """
    
    try:
        response = model.generate_content(prompt)
        raw_text = response.text.strip()
        
        if raw_text.startswith("```json"):
            raw_text = raw_text.split("```json")[1]
        if raw_text.endswith("```"):
            raw_text = raw_text.rsplit("```", 1)[0]
            
        raw_text = raw_text.strip()
        import json
        result = json.loads(raw_text)
        return result
    except Exception as e:
        print(f"[❌ EROARE AUTOGRADER] {e}")
        return {"suggestedGrade": 0, "feedback": "Eroare la evaluarea AI. Corectează manual."}

@app.post("/generate-podcast")
def generate_podcast(req: PodcastRequest):
    context_text = ""
    file_name = req.file_name
    ext = file_name.split('.')[-1].lower() if '.' in file_name else ''
    
    if ext not in ['pdf', 'docx', 'txt', 'pptx']:
        raise HTTPException(status_code=400, detail="Tip de fișier nesuportat.")
         
    try:
        response = minio_client.get_object("uniconnect-documents", file_name)
        file_content = response.read()
        response.close()
        response.release_conn()
        
        if ext == 'pdf':
            pdf_document = fitz.open(stream=file_content, filetype="pdf")
            for page_num in range(len(pdf_document)):
                context_text += pdf_document.load_page(page_num).get_text() + "\n"
            pdf_document.close()
        elif ext == 'txt':
            context_text += file_content.decode('utf-8') + "\n"
        elif ext == 'docx':
            doc = Document(io.BytesIO(file_content))
            for para in doc.paragraphs:
                context_text += para.text + "\n"
        elif ext == 'pptx':
            prs = Presentation(io.BytesIO(file_content))
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        context_text += shape.text + "\n"
    except Exception as e:
        print(f"[⚠️ AVERTISMENT] Nu am putut citi {file_name}: {e}")
        raise HTTPException(status_code=500, detail=f"Eroare citire: {str(e)}")
        
    if not context_text.strip():
        raise HTTPException(status_code=400, detail="Fișierul este gol.")
         
    context_text = context_text[:150000]
    
    prompt = f"""Ești un prezentator de podcast educațional. Scrie un script de podcast (doar discursul tău, fără indicații de scenă, direcții sau nume de personaje). 
Vorbește la persoana I, explică cele mai importante 3-4 concepte din textul următor pe un ton relaxat, antrenant și simplu de înțeles (max 3 minute de vorbire). Fii entuziast! 
Limba: Română. 

Textul de curs:
{context_text}
"""
    
    try:
        response = model.generate_content(prompt)
        script_text = response.text.strip()
        
        script_text = script_text.replace("*", "")
        script_text = script_text.replace("#", "")
        
        tts = gTTS(text=script_text, lang='ro', slow=False)
        mp3_fp = io.BytesIO()
        tts.write_to_fp(mp3_fp)
        mp3_fp.seek(0)
        
        new_file_name = f"podcast_{uuid.uuid4().hex[:8]}.mp3"
        length = mp3_fp.getbuffer().nbytes
        
        minio_client.put_object(
            "uniconnect-documents",
            new_file_name,
            mp3_fp,
            length,
            content_type="audio/mpeg"
        )
        
        return {"podcastFileName": new_file_name, "script": script_text}

    except Exception as e:
        print(f"[❌ EROARE PODCAST] {e}")
        raise HTTPException(status_code=500, detail="Eroare la generarea podcastului.")

if __name__ == "__main__":
    uvicorn.run(app, host="0.0.0.0", port=5001)